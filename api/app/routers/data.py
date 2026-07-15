from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.encoders import jsonable_encoder
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from pydantic import BaseModel, Field
from sqlalchemy import Select, func, select
from sqlalchemy.orm import Session
from sqlalchemy.sql.elements import ColumnElement

from app.auth.tokens import get_current_user
from app.authz.fields import allowed_columns
from app.authz.scopes import scope_predicate
from app.db.models import AppUser, AuditLog, Base, ColumnMeta, TableMeta
from app.db.session import get_db
from app.meta.resolver import MetadataResolutionError, ResolvedQueryPlan, resolve_query_plan

router = APIRouter(prefix="/api/data", tags=["data"])
EXPORT_ROW_CAP = 50_000


class FilterItem(BaseModel):
    col: str
    op: str
    val: Any = None


class SortItem(BaseModel):
    col: str
    dir: str = "asc"


class QueryRequest(BaseModel):
    filters: list[FilterItem] = Field(default_factory=list)
    sort: list[SortItem] = Field(default_factory=list)
    page: int = Field(default=1, ge=1)
    size: int = Field(default=50, ge=1, le=200)


class ExportRequest(BaseModel):
    filters: list[FilterItem] = Field(default_factory=list)
    sort: list[SortItem] = Field(default_factory=list)


def _role_allowed(role: str, roles: list[str]) -> bool:
    return role in roles


def _metadata_for_table(db: Session, table_code: str, user: AppUser) -> tuple[TableMeta, list[ColumnMeta]]:
    table = db.scalar(select(TableMeta).where(TableMeta.code == table_code, TableMeta.deleted_at.is_(None)))
    if table is None:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unregistered table")
    if not _role_allowed(user.role, table.read_roles):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Forbidden")

    columns = db.scalars(
        select(ColumnMeta)
        .where(ColumnMeta.table_code == table_code, ColumnMeta.deleted_at.is_(None))
        .order_by(ColumnMeta.id)
    ).all()
    return table, columns


def _filter_dicts(filters: list[FilterItem]) -> list[dict[str, Any]]:
    return [item.model_dump() for item in filters]


def _check_filter_roles(columns_by_code: dict[str, ColumnMeta], filters: list[FilterItem], role: str) -> None:
    for item in filters:
        column = columns_by_code.get(item.col)
        if column is None:
            continue
        if not _role_allowed(role, column.filter_roles):
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN, detail="Column is not filterable for role"
            )


def _check_sort_roles(columns_by_code: dict[str, ColumnMeta], sort: list[SortItem], role: str) -> None:
    for item in sort:
        column = columns_by_code.get(item.col)
        if column is None:
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unregistered sort column")
        if item.dir.lower() not in {"asc", "desc"}:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST, detail="Sort dir must be asc or desc"
            )
        if not _role_allowed(role, column.sort_roles):
            raise HTTPException(
                status_code=status.HTTP_403_FORBIDDEN, detail="Column is not sortable for role"
            )


def _resolve_plan(
    db: Session,
    table_code: str,
    column_codes: list[str],
    filters: list[FilterItem],
) -> ResolvedQueryPlan:
    try:
        return resolve_query_plan(db, table_code, column_codes, _filter_dicts(filters))
    except MetadataResolutionError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


def _resolve_sort_columns(db: Session, table_code: str, sort: list[SortItem]) -> None:
    if not sort:
        return
    try:
        resolve_query_plan(db, table_code, [item.col for item in sort], [])
    except MetadataResolutionError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc


def _physical_table(plan: ResolvedQueryPlan):
    table = Base.metadata.tables.get(plan.physical_table)
    if table is None:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unregistered physical table")
    return table


def _coerce_range(value: Any) -> tuple[Any, Any]:
    if not isinstance(value, list | tuple) or len(value) != 2:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail="Range filter requires [from, to]"
        )
    return value[0], value[1]


def _filter_expression(column: ColumnElement, operator: str, value: Any) -> ColumnElement[bool]:
    if operator == "eq":
        return column == value
    if operator == "contains":
        return column.ilike(f"%{value}%")
    if operator == "in":
        if not isinstance(value, list):
            raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="In filter requires a list")
        return column.in_(value)
    if operator == "range":
        start, end = _coerce_range(value)
        expressions: list[ColumnElement[bool]] = []
        if start is not None:
            expressions.append(column >= start)
        if end is not None:
            expressions.append(column <= end)
        if not expressions:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST, detail="Range filter needs a boundary"
            )
        expression = expressions[0]
        for extra in expressions[1:]:
            expression = expression & extra
        return expression
    if operator == "isnull":
        return column.is_(None) if bool(value) else column.is_not(None)
    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Operator not allowed")


def _apply_filters(statement: Select, table, plan: ResolvedQueryPlan, filters: list[FilterItem]) -> Select:
    for resolved, item in zip(plan.filters, filters, strict=True):
        statement = statement.where(
            _filter_expression(table.c[resolved.physical_column], resolved.operator or item.op, item.val)
        )
    return statement


def _apply_scope(statement: Select, user: AppUser, table_code: str) -> Select:
    predicate = scope_predicate(user, table_code)
    if predicate is None:
        return statement
    return statement.where(predicate)


def _apply_active(statement: Select, table) -> Select:
    if "deleted_at" in table.c:
        return statement.where(table.c.deleted_at.is_(None))
    return statement


def _apply_sort(
    statement: Select, table, columns_by_code: dict[str, ColumnMeta], sort: list[SortItem]
) -> Select:
    for item in sort:
        physical = columns_by_code[item.col].physical_column
        column = table.c[physical]
        statement = statement.order_by(column.desc() if item.dir.lower() == "desc" else column.asc())
    if not sort and "id" in table.c:
        statement = statement.order_by(table.c.id.asc())
    return statement


def _build_context(
    db: Session,
    table_code: str,
    user: AppUser,
    filters: list[FilterItem],
    sort: list[SortItem],
    *,
    export: bool = False,
) -> tuple[ResolvedQueryPlan, Any, dict[str, ColumnMeta]]:
    _, columns = _metadata_for_table(db, table_code, user)
    columns_by_code = {column.col_code: column for column in columns}
    readable = allowed_columns(user.role, table_code, db)
    if export:
        column_codes = [
            column.col_code
            for column in columns
            if column.col_code in readable and _role_allowed(user.role, column.export_roles)
        ]
    else:
        column_codes = [column.col_code for column in columns if column.col_code in readable]

    _check_filter_roles(columns_by_code, filters, user.role)
    _resolve_sort_columns(db, table_code, sort)
    _check_sort_roles(columns_by_code, sort, user.role)
    plan = _resolve_plan(db, table_code, column_codes, filters)
    return plan, _physical_table(plan), columns_by_code


def _base_statement(table, plan: ResolvedQueryPlan):
    selected = [table.c[column.physical_column].label(column.code) for column in plan.columns]
    return select(*selected).select_from(table)


def _count_statement(table):
    return select(func.count()).select_from(table)


def _apply_query_pipeline(
    statement: Select,
    table,
    plan: ResolvedQueryPlan,
    filters: list[FilterItem],
    user: AppUser,
) -> Select:
    statement = _apply_active(statement, table)
    statement = _apply_filters(statement, table, plan, filters)
    return _apply_scope(statement, user, plan.table_code)


@router.post("/{table_code}/query")
def query_data(
    table_code: str,
    payload: QueryRequest,
    current_user: AppUser = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> dict[str, Any]:
    plan, table, columns_by_code = _build_context(db, table_code, current_user, payload.filters, payload.sort)
    count_statement = _apply_query_pipeline(
        _count_statement(table), table, plan, payload.filters, current_user
    )
    total = db.scalar(count_statement) or 0

    rows_statement = _apply_query_pipeline(
        _base_statement(table, plan), table, plan, payload.filters, current_user
    )
    rows_statement = _apply_sort(rows_statement, table, columns_by_code, payload.sort)
    rows_statement = rows_statement.offset((payload.page - 1) * payload.size).limit(payload.size)
    rows = [dict(row) for row in db.execute(rows_statement).mappings().all()]
    db.add(
        AuditLog(
            actor=current_user.id,
            action="query",
            table_code=plan.table_code,
            filters={"filters": _filter_dicts(payload.filters)},
            row_count=total,
        )
    )
    db.commit()
    return {"rows": jsonable_encoder(rows), "total": total, "page": payload.page, "size": payload.size}


def _write_cell(cell: Cell, value: Any, column_type: str) -> None:
    if value is None:
        cell.value = None
        return
    if column_type in {"text", "ym"}:
        cell.value = str(value)
        cell.number_format = "@"
        return
    if isinstance(value, Decimal):
        cell.value = float(value)
        return
    if isinstance(value, date | datetime):
        cell.value = value
        cell.number_format = "yyyy-mm-dd"
        return
    cell.value = value


def _workbook_bytes(
    rows: list[dict[str, Any]], plan: ResolvedQueryPlan, columns_by_code: dict[str, ColumnMeta]
) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([columns_by_code[column.code].label for column in plan.columns])
    for row_index, row in enumerate(rows, start=2):
        for column_index, column in enumerate(plan.columns, start=1):
            _write_cell(sheet.cell(row=row_index, column=column_index), row.get(column.code), column.type)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@router.post("/{table_code}/export")
def export_data(
    table_code: str,
    payload: ExportRequest,
    current_user: AppUser = Depends(get_current_user),
    db: Session = Depends(get_db),
) -> StreamingResponse:
    plan, table, columns_by_code = _build_context(
        db, table_code, current_user, payload.filters, payload.sort, export=True
    )
    count_statement = _apply_query_pipeline(
        _count_statement(table), table, plan, payload.filters, current_user
    )
    row_count = db.scalar(count_statement) or 0
    if row_count > EXPORT_ROW_CAP:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Export exceeds {EXPORT_ROW_CAP} rows; narrow the filters",
        )

    rows_statement = _apply_query_pipeline(
        _base_statement(table, plan), table, plan, payload.filters, current_user
    )
    rows_statement = _apply_sort(rows_statement, table, columns_by_code, payload.sort).limit(EXPORT_ROW_CAP)
    rows = [dict(row) for row in db.execute(rows_statement).mappings().all()]
    db.add(
        AuditLog(
            actor=current_user.id,
            action="export",
            table_code=plan.table_code,
            filters={"filters": _filter_dicts(payload.filters)},
            row_count=row_count,
        )
    )
    db.commit()

    output = _workbook_bytes(rows, plan, columns_by_code)
    timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
    filename = f"{plan.physical_table}_{timestamp}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
